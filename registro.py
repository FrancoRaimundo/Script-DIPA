#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Registro coloreado de carga (DIPA / MAIBA)
------------------------------------------
Mantiene un archivo .xlsx con una fila por participante y la casilla del DNI
PINTADA segun el estado, automatizando lo que el operador hacia a mano:

  - VERDE   -> CARGADO   (se cargo bien en MAIBA)
  - ROJO    -> ERROR / FALTA / INCOMPLETO (no se pudo / faltan pdf-jpeg / datos)
  - AMARILLO-> LISTO / PENDIENTE (validado y a la espera de cargar)

Se usa como fuente de verdad reanudable: el cargador lee que DNIs ya estan en
VERDE (CARGADO) para no repetirlos, y actualiza cada fila al terminar.
"""

from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import PatternFill, Font
except ImportError:
    raise SystemExit("Falta openpyxl. Instalalo con:  pip install openpyxl")


VERDE = PatternFill("solid", fgColor="C6EFCE")
ROJO = PatternFill("solid", fgColor="FFC7CE")
AMARILLO = PatternFill("solid", fgColor="FFEB9C")

# Estado -> color de la casilla
COLOR = {
    "CARGADO":  VERDE,
    "LISTO":    AMARILLO,
    "PENDIENTE": AMARILLO,
    "SALTEADO": AMARILLO,
    "ERROR":    ROJO,
    "FALTA":    ROJO,
    "INCOMPLETO": ROJO,
    "REVISAR":  ROJO,
}

ENCABEZADOS = ["DNI", "Apellido", "Nombre", "CUIL", "Estado", "Detalle"]
COL_DNI = 1
COL_ESTADO = 5


def _color(estado):
    return COLOR.get(str(estado).upper(), AMARILLO)


def escribir(ruta, filas):
    """Crea (o reemplaza) el registro con la lista de filas. Cada fila es un
    dict con claves: dni, apellido, nombre, cuil, estado, detalle."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    ws.append(ENCABEZADOS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append([f.get("dni", ""), f.get("apellido", ""), f.get("nombre", ""),
                   f.get("cuil", ""), f.get("estado", ""), f.get("detalle", "")])
        fila = ws.max_row
        fill = _color(f.get("estado", ""))
        ws.cell(fila, COL_DNI).fill = fill
        ws.cell(fila, COL_ESTADO).fill = fill
    # Anchos comodos
    for col, ancho in zip("ABCDEF", (12, 18, 18, 14, 12, 50)):
        ws.column_dimensions[col].width = ancho
    wb.save(ruta)


def leer_cargados(ruta):
    """Devuelve el set de DNIs cuyo Estado es CARGADO (para reanudar)."""
    p = Path(ruta)
    if not p.exists():
        return set()
    wb = load_workbook(p)
    ws = wb.active
    hechos = set()
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila and str(fila[COL_ESTADO - 1]).upper() == "CARGADO":
            hechos.add(str(fila[COL_DNI - 1]))
    return hechos


def actualizar(ruta, dni, estado, detalle=""):
    """Cambia el estado/color de la fila de un DNI. Si no existe, la agrega."""
    p = Path(ruta)
    if not p.exists():
        escribir(ruta, [{"dni": dni, "estado": estado, "detalle": detalle}])
        return
    wb = load_workbook(p)
    ws = wb.active
    objetivo = None
    for fila in range(2, ws.max_row + 1):
        if str(ws.cell(fila, COL_DNI).value) == str(dni):
            objetivo = fila
            break
    if objetivo is None:
        ws.append([dni, "", "", "", estado, detalle])
        objetivo = ws.max_row
    else:
        ws.cell(objetivo, COL_ESTADO).value = estado
        ws.cell(objetivo, 6).value = detalle
    fill = _color(estado)
    ws.cell(objetivo, COL_DNI).fill = fill
    ws.cell(objetivo, COL_ESTADO).fill = fill
    wb.save(ruta)
